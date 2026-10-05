import logging
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import Config

logger = logging.getLogger("attendance_system.excel_service")


class ExcelService:
    """Service to handle Excel synchronization, exports, and backup records."""

    @staticmethod
    def _apply_table_styling(ws, header_title="ATTENDANCE SYSTEM REPORT"):
        """Apply modern, professional typography and styling to an openpyxl worksheet."""
        # Brand colors
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate 800
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")     # Blue 500
        sub_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10, color="1E293B")
        bold_font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Style Header Row (Row 4 is typical column header)
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.fill = sub_fill
            cell.font = sub_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Style Data Rows
        for row_idx in range(5, ws.max_row + 1):
            is_even = (row_idx % 2 == 0)
            row_fill = PatternFill(start_color="F8FAFC" if is_even else "FFFFFF", fill_type="solid")
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.fill = row_fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", horizontal="left")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if cell.row < 4:  # Title banner
                    continue
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def sync_to_master(cls, participant_data: dict, attendance_data: dict = None):
        """
        Synchronize or update participant record in Master_Attendance_Record.xlsx.
        Keeps an organized, updated backup spreadsheet.
        """
        Config.EXCEL_MASTER_FILE.parent.mkdir(parents=True, exist_ok=True)
        file_path = Config.EXCEL_MASTER_FILE

        if not file_path.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Master Attendance Record"
            
            # Title Banner
            ws.merge_cells("A1:J2")
            title_cell = ws["A1"]
            title_cell.value = "REAL-TIME ID VERIFICATION & ATTENDANCE SYSTEM - MASTER RECORD"
            title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
            title_cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")

            # Column Headers
            headers = [
                "Participant ID (UUID)",
                "Full Name",
                "Registration Number",
                "Email Address",
                "Phone Number",
                "Event / Classroom",
                "Category",
                "Registration Time",
                "Attendance Status",
                "Entry Time",
                "Exit Time"
            ]
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=4, column=col_idx, value=header)
            wb.save(file_path)

        # Open and update/append
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb["Master Attendance Record"]
            reg_num = str(participant_data.get("registration_number", "")).strip()

            target_row = None
            # Search for existing registration number in column 3
            for row in range(5, ws.max_row + 1):
                cell_val = str(ws.cell(row=row, column=3).value or "").strip()
                if cell_val == reg_num:
                    target_row = row
                    break

            if target_row is None:
                target_row = ws.max_row + 1

            # Populate row
            ws.cell(row=target_row, column=1, value=str(participant_data.get("participant_uuid", "")))
            ws.cell(row=target_row, column=2, value=str(participant_data.get("name", "")))
            ws.cell(row=target_row, column=3, value=reg_num)
            ws.cell(row=target_row, column=4, value=str(participant_data.get("email", "")))
            ws.cell(row=target_row, column=5, value=str(participant_data.get("phone", "")))
            ws.cell(row=target_row, column=6, value=str(participant_data.get("event_name", "")))
            ws.cell(row=target_row, column=7, value=str(participant_data.get("category", "Participant")))
            ws.cell(row=target_row, column=8, value=str(participant_data.get("created_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))))

            # Attendance fields
            status = "Absent"
            entry_time = ""
            exit_time = ""

            if attendance_data:
                status = attendance_data.get("status", "Present")
                entry_time = str(attendance_data.get("entry_time", ""))
                exit_time = str(attendance_data.get("exit_time", "") or "")
            elif ws.cell(row=target_row, column=9).value:
                status = ws.cell(row=target_row, column=9).value
                entry_time = ws.cell(row=target_row, column=10).value or ""
                exit_time = ws.cell(row=target_row, column=11).value or ""

            ws.cell(row=target_row, column=9, value=status)
            ws.cell(row=target_row, column=10, value=entry_time)
            ws.cell(row=target_row, column=11, value=exit_time)

            cls._apply_table_styling(ws)
            wb.save(file_path)
            logger.info(f"Synchronized participant {reg_num} to Master Excel record.")
        except Exception as e:
            logger.error(f"Error syncing to master Excel record: {e}")

    @classmethod
    def generate_attendance_report(cls, records: list, event_name: str = "All Events") -> str:
        """
        Generate a downloadable, styled Excel report from attendance query results.
        Returns the saved file path.
        """
        Config.EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        sanitized_event = "".join(c for c in event_name if c.isalnum() or c in (' ', '_', '-')).strip()
        filename = f"Attendance_Report_{sanitized_event}_{timestamp}.xlsx"
        filepath = Config.EXPORTS_DIR / filename

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Summary"

        # Title Banner
        ws.merge_cells("A1:H2")
        title_cell = ws["A1"]
        title_cell.value = f"ATTENDANCE REPORT - {event_name.upper()}"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Records: {len(records)}")
        ws.cell(row=3, column=1).font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

        headers = [
            "Sl No",
            "Participant Name",
            "Registration No",
            "Event Name",
            "Attendance Status",
            "Entry Time",
            "Exit Time",
            "Verification Method"
        ]
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=4, column=col_idx, value=header)

        for row_idx, rec in enumerate(records, 5):
            ws.cell(row=row_idx, column=1, value=row_idx - 4)
            ws.cell(row=row_idx, column=2, value=rec.get("name", ""))
            ws.cell(row=row_idx, column=3, value=rec.get("registration_number", ""))
            ws.cell(row=row_idx, column=4, value=rec.get("event_name", ""))
            ws.cell(row=row_idx, column=5, value=rec.get("status", "Present"))
            ws.cell(row=row_idx, column=6, value=str(rec.get("entry_time", "")))
            ws.cell(row=row_idx, column=7, value=str(rec.get("exit_time", "") or "-"))
            ws.cell(row=row_idx, column=8, value=rec.get("verification_method", "QR_SCAN"))

        cls._apply_table_styling(ws)
        wb.save(filepath)
        logger.info(f"Generated attendance Excel report: {filepath}")
        return str(filepath)
